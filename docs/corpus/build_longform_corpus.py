# -*- coding: utf-8 -*-
"""Fetch official long-form China-story articles and compile into a Word file."""
from __future__ import annotations

import json
import time
from pathlib import Path

import requests
import trafilatura
from docx import Document
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

ROOT = Path(__file__).resolve().parent
CATALOG = ROOT / "longform_catalog.json"
OUT_DOCX = ROOT / "讲好中国故事_官媒长文语料库_约100篇.docx"
OUT_JSONL = ROOT / "longform_fetched.jsonl"
HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
        "AppleWebKit/537.36 (KHTML, like Gecko) "
        "Chrome/122.0.0.0 Safari/537.36"
    ),
    "Accept-Language": "zh-CN,zh;q=0.9,en;q=0.8",
}


def set_run_font(run, name_cn="宋体", name_en="Calibri", size=11, bold=False, color=None):
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = name_en
    r = run._element
    rPr = r.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:eastAsia"), name_cn)
    if color is not None:
        run.font.color.rgb = color


def add_heading_cn(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        set_run_font(run, size=16 if level == 1 else 13, bold=True, name_cn="黑体")
    return p


def fetch_one(url: str, timeout: int = 25) -> tuple[str, str, str]:
    """Return (status, title, text)."""
    try:
        downloaded = trafilatura.fetch_url(url)
        if not downloaded:
            resp = requests.get(url, headers=HEADERS, timeout=timeout)
            resp.raise_for_status()
            # prefer apparent encoding for Chinese pages
            if not resp.encoding or resp.encoding.lower() in ("iso-8859-1", "ascii"):
                resp.encoding = resp.apparent_encoding or "utf-8"
            downloaded = resp.text
        meta = trafilatura.bare_extraction(
            downloaded,
            include_comments=False,
            include_tables=False,
            favor_precision=True,
            url=url,
        )
        if meta is None:
            text = trafilatura.extract(downloaded, favor_precision=True, url=url) or ""
            title = ""
        else:
            # trafilatura Document-like object or dict depending on version
            if isinstance(meta, dict):
                title = (meta.get("title") or "").strip()
                text = (meta.get("text") or "").strip()
            else:
                title = (getattr(meta, "title", None) or "").strip()
                text = (getattr(meta, "text", None) or "").strip()
        if not text or len(text) < 200:
            return "thin", title, text
        return "ok", title, text
    except Exception as e:  # noqa: BLE001
        return f"error:{type(e).__name__}:{e}", "", ""


def build_doc(items: list[dict], results: list[dict]) -> Document:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

    add_heading_cn(doc, "讲好中国故事 · 官媒长文 / 新闻通稿语料库", 0)
    p = doc.add_paragraph()
    set_run_font(
        p.add_run(
            "用途：支撑长文（news / feature）体裁开发与风格参照。"
            "收录近年官方媒体讲述中国的通稿/长通讯/英文外宣长文，优先对齐现有主题："
            "光伏羊/牧光互补、沙戈荒光伏治沙、乡村振兴、绿色低碳、生态修复、共同富裕、"
            "中国式现代化地方样本、人文交流与一带一路民生等。"
        ),
        size=10.5,
    )
    p = doc.add_paragraph()
    set_run_font(
        p.add_run(
            "结构：A 中国互联网官媒约30条；B 外网官方英文站点约70条"
            "（China Daily / CGTN / Xinhua English / People's Daily Online / SCIO / China.org.cn 等，"
            "亦为 Facebook、Instagram、X 等官方账号常转发的长文原文页面）。"
            "正文为抓取时页面可提取文本；若标注「抓取失败/正文偏短」，请按 URL 人工核验。"
            "版权归原媒体所有，仅供团队内部研发参照，勿对外再发布。"
        ),
        size=10,
        color=RGBColor(0x55, 0x55, 0x55),
    )

    by_id = {r["id"]: r for r in results}
    sections = [
        ("A", "中国互联网官方媒体长文 / 通稿（约30条）"),
        ("B", "外网官方媒体长文（约70条，英文站点为主）"),
    ]
    for sec_code, sec_title in sections:
        add_heading_cn(doc, sec_title, 1)
        sec_items = [it for it in items if it["section"] == sec_code]
        for it in sec_items:
            r = by_id.get(it["id"], {})
            status = r.get("status", "missing")
            title = r.get("fetched_title") or it["title"]
            text = r.get("text") or ""
            add_heading_cn(doc, f"{it['id']}. {title}", 2)
            meta = doc.add_paragraph()
            set_run_font(
                meta.add_run(
                    f"来源：{it['source']}｜日期：{it.get('date', '')}｜"
                    f"主题：{it.get('tags', '')}｜状态：{status}\nURL：{it['url']}"
                ),
                size=9,
                color=RGBColor(0x33, 0x33, 0x99),
            )
            if text:
                for para in text.split("\n"):
                    para = para.strip()
                    if not para:
                        continue
                    pp = doc.add_paragraph()
                    set_run_font(pp.add_run(para), size=11)
            else:
                pp = doc.add_paragraph()
                set_run_font(
                    pp.add_run("【未能自动提取全文，请打开上方 URL 人工复制正文。】"),
                    size=10,
                    color=RGBColor(0x99, 0x33, 0x33),
                )
            doc.add_paragraph()
    return doc


def main():
    items = json.loads(CATALOG.read_text(encoding="utf-8"))
    results = []
    print(f"catalog size: {len(items)}")
    with OUT_JSONL.open("w", encoding="utf-8") as fout:
        for i, it in enumerate(items, 1):
            print(f"[{i}/{len(items)}] {it['id']} {it['url'][:80]}...")
            status, title, text = fetch_one(it["url"])
            row = {
                "id": it["id"],
                "section": it["section"],
                "url": it["url"],
                "catalog_title": it["title"],
                "fetched_title": title,
                "status": status,
                "text": text,
                "chars": len(text),
            }
            results.append(row)
            fout.write(json.dumps(row, ensure_ascii=False) + "\n")
            fout.flush()
            time.sleep(0.6)

    ok = sum(1 for r in results if r["status"] == "ok")
    thin = sum(1 for r in results if r["status"] == "thin")
    err = sum(1 for r in results if str(r["status"]).startswith("error"))
    print(f"done ok={ok} thin={thin} err={err}")

    doc = build_doc(items, results)
    doc.save(OUT_DOCX)
    print(f"wrote {OUT_DOCX}")

    # summary sidecar
    summary = {
        "total": len(items),
        "ok": ok,
        "thin": thin,
        "err": err,
        "section_A": sum(1 for it in items if it["section"] == "A"),
        "section_B": sum(1 for it in items if it["section"] == "B"),
        "docx": str(OUT_DOCX),
    }
    (ROOT / "longform_fetch_summary.json").write_text(
        json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8"
    )


if __name__ == "__main__":
    main()
