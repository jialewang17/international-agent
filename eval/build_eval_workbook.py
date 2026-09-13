"""Build multi-sheet Excel workbook for large-scale skill eval (v0.3 skeleton)."""

from __future__ import annotations

import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent
CASES_CSV = ROOT / "skill_eval_cases_v0.3_40.csv"
OUT_XLSX = ROOT / "skill_eval_v0.3_workbook.xlsx"

HEADER_FILL = PatternFill("solid", fgColor="C41E3A")
HEADER_FONT = Font(color="FFFFFF", bold=True)
WRAP = Alignment(wrap_text=True, vertical="top")


def _style_header(ws, ncols: int) -> None:
    for col in range(1, ncols + 1):
        cell = ws.cell(1, col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = WRAP


def _autosize(ws, max_width: int = 48) -> None:
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        length = 0
        for cell in col[:40]:
            if cell.value:
                length = max(length, min(len(str(cell.value)), max_width))
        ws.column_dimensions[letter].width = max(12, length + 2)


def main() -> None:
    rows = list(csv.DictReader(CASES_CSV.open(encoding="utf-8-sig")))
    wb = Workbook()

    # --- 01 评测明细 ---
    ws1 = wb.active
    ws1.title = "01_评测明细"
    h1 = [
        "测试ID",
        "组别",
        "能力线",
        "调用工具",
        "Skill版本",
        "测试日期",
        "执行人",
        "evidence条数",
        "门控是否通过(是/否)",
        "D1主题一致性(1-5)",
        "D2事实准确性(1-5)",
        "D3平台适配(1-5)",
        "D4语言自然度(1-5)",
        "D5国际传播方式(1-5)",
        "D6素材运用(1-5)",
        "D7风险控制(1-5)",
        "总分(35)",
        "是否达标(是/否)",
        "主要问题",
        "Agent完整输出(路径)",
        "备注",
        "扩样依据",
    ]
    ws1.append(h1)
    for r in rows:
        ws1.append(
            [
                r["测试ID"],
                r["组别"],
                r["能力线"],
                r["调用工具"],
                "v0.2+gate",
                "",
                "",
                "",
                "",
                "",
                "",
                "",
                "",
                "",
                "",
                "",
                "",
                "",
                "",
                f"eval/outputs/{r['测试ID']}.json",
                "待R4跑测填分",
                r.get("扩样依据", ""),
            ]
        )
    _style_header(ws1, len(h1))
    _autosize(ws1)

    # --- 02 用例与Prompt ---
    ws2 = wb.create_sheet("02_用例与Prompt")
    h2 = ["测试ID", "组别", "能力线", "调用工具", "用户输入Prompt", "预期行为", "扩样依据"]
    ws2.append(h2)
    for r in rows:
        ws2.append([r[k] for k in ["测试ID", "组别", "能力线", "调用工具", "用户输入Prompt（复制到CLI）", "预期行为（门控/内容）", "扩样依据"]])
    _style_header(ws2, len(h2))
    _autosize(ws2, 60)

    # --- 03 Changelog ---
    ws3 = wb.create_sheet("03_迭代Changelog")
    h3 = [
        "迭代轮次",
        "日期",
        "改动文件",
        "版本",
        "本次只改一处",
        "修改位置",
        "改了什么(摘要)",
        "优化目标",
        "改前相关均分/达标",
        "改后相关均分/达标",
        "是否保留",
        "评审人确认",
        "备注",
    ]
    ws3.append(h3)
    ws3.append(
        [
            1,
            "2026-08-22",
            "skills/intl-comm.md",
            "v0.1",
            "基线",
            "",
            "建立15题评测基线",
            "T01-T15",
            "",
            "均分26.0/35",
            "是",
            "",
            "桌面xlsx",
        ]
    )
    ws3.append(
        [
            2,
            "2026-08-22",
            "skills/intl-comm.md",
            "v0.1→v0.2",
            "证据与风控分层",
            "§v0.2 A/B/C/D",
            "硬性/演示期+语气边界",
            "T05/T11/T14等",
            "",
            "待记",
            "是",
            "已通过",
            "",
        ]
    )
    ws3.append(
        [
            3,
            "2026-08-22",
            "kb_local+story_post_gen",
            "v0.2+gate",
            "主题-素材门控",
            "check_theme_evidence_alignment",
            "库外/缺条拒编",
            "T09/T10/T14",
            "门控漏拦",
            "T09/10/14拒编",
            "是",
            "已通过",
            "",
        ]
    )
    ws3.append(
        [
            4,
            "2026-09-06",
            "评测体系",
            "v0.3-expand",
            "扩样至40题+反馈表",
            "eval/",
            "新增T16-T40与多表workbook",
            "老师要求大规模+反馈",
            "15题",
            "待跑",
            "待确认",
            "等你确认后执行R4跑测",
            "大改skill前必问",
        ]
    )
    _style_header(ws3, len(h3))
    _autosize(ws3)

    # --- 04 提示词快照 ---
    ws4 = wb.create_sheet("04_提示词快照")
    h4 = ["轮次", "Skill版本", "文件路径", "用途", "摘要/说明", "是否变更"]
    ws4.append(h4)
    snaps = [
        ("R4基线", "v0.2", "skills/intl-comm.md", "团队叙事+证据风控", "硬性/演示期分层；语气边界；自检四问", "本轮暂不改"),
        ("R4基线", "v0.2", "skills/china-story-active.md", "主动传播流程", "引用intl-comm v0.2", "本轮暂不改"),
        ("R4基线", "—", "prompt/story_post_prompt.txt", "成稿LLM", "事实须来自evidence", "本轮暂不改"),
        ("R4基线", "—", "prompt/reply_generation_prompt.txt", "回复LLM", "叙事向回复", "本轮暂不改"),
        ("R4基线", "—", "prompt/absa_topic_prompt.txt", "主题识别", "ABSA", "本轮暂不改"),
        ("R4基线", "v0.2+gate", "tools/kb_local.py", "对齐门控", "off_topic / subtopic_unsupported", "本轮暂不改(除非你批准扩标记)"),
    ]
    for s in snaps:
        ws4.append(list(s))
    _style_header(ws4, len(h4))
    _autosize(ws4)

    # --- 05 修改前后对比 ---
    ws5 = wb.create_sheet("05_修改前后对比")
    h5 = [
        "测试ID",
        "对比轮次",
        "改前Skill/代码",
        "改后Skill/代码",
        "改前门控",
        "改后门控",
        "改前总分",
        "改后总分",
        "改前主要问题",
        "改后主要问题",
        "改前输出路径",
        "改后输出路径",
        "结论",
    ]
    ws5.append(h5)
    ws5.append(
        [
            "T09",
            "v0.1→v0.2+gate",
            "仅空evidence门控",
            "+主题对齐门控",
            "否(仍生成)",
            "是(拒编)",
            15,
            "≈31",
            "硬凑冬奥",
            "off_topic_foreign",
            "eval/outputs历史",
            "eval/outputs/T09.json",
            "门控有效",
        ]
    )
    ws5.append(
        [
            "T11",
            "v0.1→v0.2 skill",
            "无强制反嘲讽",
            "§C语气边界",
            "是",
            "是(仍生成嘲讽)",
            24,
            "待R4重测",
            "嘲讽美国",
            "待填",
            "历史",
            "eval/outputs/T11.json",
            "Skill alone不足→R5候选",
        ]
    )
    _style_header(ws5, len(h5))
    _autosize(ws5)

    # --- 06 反馈自修正日志 ---
    ws6 = wb.create_sheet("06_反馈自修正日志")
    h6 = [
        "日志ID",
        "日期",
        "失败测试ID",
        "诊断标签",
        "对应维度",
        "建议动作(只改一处)",
        "是否采纳(是/否/改)",
        "你的意见摘要",
        "落地文件",
        "重测ID列表",
        "重测结果摘要",
        "是否进入下一轮",
    ]
    ws6.append(h6)
    ws6.append(
        [
            "F001",
            "2026-08-22",
            "T09,T10,T14",
            "门控漏拦/子话题缺条",
            "D2/D6",
            "后端对齐门控",
            "是",
            "同意改后端",
            "kb_local/story_post_gen",
            "T09,T10,T14",
            "均正确STOP",
            "否(已完成)",
        ]
    )
    ws6.append(
        [
            "F002",
            "待R4",
            "待填",
            "待诊断",
            "",
            "R4跑完后自动/人工写入候选",
            "待问你",
            "",
            "",
            "",
            "",
            "是",
        ]
    )
    _style_header(ws6, len(h6))
    _autosize(ws6)

    # --- 07 汇总看板 ---
    ws7 = wb.create_sheet("07_汇总看板")
    h7 = [
        "Skill版本",
        "评测日期",
        "用例总数",
        "已完成数",
        "门控通过数",
        "A达标",
        "B达标",
        "C达标",
        "D1均分",
        "D2均分",
        "D3均分",
        "D4均分",
        "D5均分",
        "D6均分",
        "D7均分",
        "总均分",
        "最低维度",
        "本轮重点问题",
        "下一轮计划",
        "评审确认",
    ]
    ws7.append(h7)
    ws7.append(
        [
            "v0.1",
            "2026-08-22",
            15,
            15,
            13,
            "6/7",
            "1/3",
            "3/5",
            4.0,
            3.4,
            3.8,
            3.8,
            4.1,
            3.3,
            3.5,
            26.0,
            "D6",
            "门控漏拦+库外编造+T11",
            "skill v0.2+门控",
            "已完成",
        ]
    )
    ws7.append(
        [
            "v0.2+gate / R4待跑",
            "",
            40,
            0,
            "",
            "",
            "",
            "",
            "",
            "",
            "",
            "",
            "",
            "",
            "",
            "",
            "",
            "等你确认后跑T01-T40",
            "按失败Top1改一处",
            "待你确认",
        ]
    )
    _style_header(ws7, len(h7))
    _autosize(ws7)

    # --- 00 使用说明 ---
    ws0 = wb.create_sheet("00_使用说明", 0)
    ws0["A1"] = "大规模评测工作簿 v0.3"
    ws0["A1"].font = Font(bold=True, size=14)
    notes = [
        "",
        "1. 对标桌面 skill_eval_v0.1_filled.xlsx 的明细列，并增加 Prompt/Changelog/对比/反馈/汇总。",
        "2. R4：仅跑当前 v0.2+gate 基线，不改 skill（除非你另批）。",
        "3. 大调整流程：写入 06_反馈自修正日志 → 问你 → 改一处 → 填 05 对比 → 更新 07。",
        "4. JSON 输出默认 eval/outputs/Txx.json；合并包可另存。",
        "5. 方案全文见 eval/LARGE_EVAL_PLAN.md",
    ]
    for i, line in enumerate(notes, start=2):
        ws0[f"A{i}"] = line
    ws0.column_dimensions["A"].width = 100

    wb.save(OUT_XLSX)
    print(f"wrote {OUT_XLSX} with {len(rows)} cases")


if __name__ == "__main__":
    main()
