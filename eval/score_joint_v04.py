"""Score joint v0.4: user materials + Q-5W/audience/authority."""

from __future__ import annotations

import json
import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "eval" / "outputs_joint_v04"
WB1 = ROOT / "eval" / "skill_eval_v0.4_joint.xlsx"
WB2 = Path.home() / "Desktop" / "skill_eval_v0.4_joint.xlsx"

RED = PatternFill("solid", fgColor="C41E3A")
WHITE = Font(color="FFFFFF", bold=True)
WRAP = Alignment(wrap_text=True, vertical="top")

import importlib.util

_spec = importlib.util.spec_from_file_location(
    "score_quality_v032", ROOT / "eval" / "score_quality_v032.py"
)
_mod = importlib.util.module_from_spec(_spec)
assert _spec.loader is not None
_spec.loader.exec_module(_mod)
score_quality = _mod.score_one


def score_joint(data: dict) -> dict:
    meta = data.get("_meta") or {}
    cid = meta.get("id", "?")
    expect_block = bool(meta.get("expect_block"))
    ev = data.get("evidence_used") or []
    n_user = sum(1 for e in ev if e.get("source_type") == "用户上传")
    n_local = sum(1 for e in ev if e.get("source_type") != "用户上传")
    has_err = bool(data.get("error")) and not (data.get("post") or "").strip()

    # gate cases
    if expect_block:
        ok = has_err
        return {
            "id": cid,
            "style": meta.get("prompt_style"),
            "n_user": n_user,
            "n_local": n_local,
            "q_5w": "-",
            "q_audience": "-",
            "q_authority": "-",
            "q_total": "-",
            "src_ok": "是" if ok else "否",
            "pass": "是" if ok else "否",
            "problem": "正确拒编" if ok else "应拒仍生成",
            "prompt": meta.get("full_prompt"),
            "expect": meta.get("expect"),
            "excerpt": (data.get("error") or "")[:180],
        }

    q = score_quality(data)
    post = (data.get("post") or "").lower()
    notes = " ".join(str(x) for x in (data.get("evidence_notes") or [])).lower()
    fw = json.dumps(data.get("five_w") or {}, ensure_ascii=False).lower()
    blob = f"{post}\n{notes}\n{fw}"
    user_statements = [e.get("statement", "") for e in ev if e.get("source_type") == "用户上传"]

    src_ok = True
    detail = []
    if meta.get("has_materials"):
        if n_user < 1:
            src_ok = False
            detail.append("缺用户上传论据")
        else:
            used = False
            # notes 明确映射 Evidence#
            if re.search(r"evidence\s*#?\s*\d", notes):
                used = True
            # 用户资料里的英文词进入正文
            for st in user_statements:
                for t in re.findall(r"[A-Za-z]{4,}", st):
                    if t.lower() in blob:
                        used = True
                        break
            # 中文资料常见译法锚点（用户事实英文化后的可检信号）
            anchors = [
                (["焊接", "打磨", "机车", "工坊", "体验团"], ["weld", "grind", "workshop", "motorcycle", "bike", "safety", "experience"]),
                (["饺子", "春联", "剪纸", "市集"], ["dumpling", "fair", "paper", "couplet", "community", "market"]),
                (["采青", "杀青", "茶园", "海拔"], ["dawn", "tea", "garden", "altitude", "leaf", "kill-green", "workshop"]),
                (["大盘鸡", "拉面", "皮带面", "乌鲁木齐"], ["chicken", "noodle", "urumqi", "laghman", "pulled"]),
                (["盖碗", "观色闻香"], ["gaiwan", "tea", "aroma", "weekend"]),
                (["崇礼", "夜滑", "缆车", "雪道"], ["chongli", "ski", "night", "gondola", "slope", "olympic"]),
                (["窗花", "刻刀", "红纸"], ["paper-cut", "papercut", "knife", "workshop", "fold"]),
                (["汉字", "节气", "书签"], ["character", "solar", "bookmark", "booth", "volunteer"]),
            ]
            for cn_keys, en_keys in anchors:
                if any(k in "".join(user_statements) for k in cn_keys):
                    if any(k in blob for k in en_keys):
                        used = True
                        break
            if not used:
                src_ok = False
                detail.append("正文/notes未体现用户事实")
            else:
                detail.append("用户事实已用于成稿")
    if q["pass"] != "是":
        detail.append(q["problem"])

    passed = q["pass"] == "是" and src_ok
    return {
        "id": cid,
        "style": meta.get("prompt_style"),
        "n_user": n_user,
        "n_local": n_local,
        "q_5w": q["q_5w"],
        "q_audience": q["q_audience"],
        "q_authority": q["q_authority"],
        "q_total": q["q_total"],
        "src_ok": "是" if src_ok else "否",
        "pass": "是" if passed else "否",
        "problem": "；".join(detail) or "联测达标",
        "prompt": meta.get("full_prompt"),
        "expect": meta.get("expect"),
        "excerpt": q.get("excerpt") or "",
    }


def main():
    rows = []
    for p in sorted(OUT.glob("J*.json")):
        rows.append(score_joint(json.loads(p.read_text(encoding="utf-8"))))

    wb = Workbook()
    ws = wb.active
    ws.title = "01_联测明细"
    headers = [
        "测试ID",
        "提示词形态",
        "提示词",
        "用户要求",
        "用户论据条数",
        "本地论据条数",
        "用户事实进正文",
        "Q-5W",
        "Q-受众",
        "Q-权威",
        "质量合计",
        "是否达标",
        "问题",
        "摘要",
    ]
    ws.append(headers)
    for r in rows:
        ws.append(
            [
                r["id"],
                r["style"],
                r["prompt"],
                r["expect"],
                r["n_user"],
                r["n_local"],
                r["src_ok"],
                r["q_5w"],
                r["q_audience"],
                r["q_authority"],
                r["q_total"],
                r["pass"],
                r["problem"],
                r["excerpt"],
            ]
        )
    for c in range(1, len(headers) + 1):
        ws.cell(1, c).fill = RED
        ws.cell(1, c).font = WHITE
        ws.cell(1, c).alignment = WRAP
    for col in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(col)].width = 16
    ws.column_dimensions["C"].width = 40

    ws2 = wb.create_sheet("02_说明")
    ws2.append(["项", "内容"])
    for row in [
        ("目的", "把用户上传论据与 v0.3.2 的5W/受众/权威标准打通联测"),
        ("顺序", "先完成Q40长短稳定性，再跑本表J01-J12"),
        ("达标", "应成稿题：质量维达标 且 用户事实进入evidence/正文；拒编题：正确STOP"),
        ("Skill", "evidence-user-materials.md + china-story-post v0.3.2 同时注入"),
    ]:
        ws2.append(list(row))
    ws2.cell(1, 1).fill = RED
    ws2.cell(1, 1).font = WHITE
    ws2.cell(1, 2).fill = RED
    ws2.cell(1, 2).font = WHITE

    passed = sum(1 for r in rows if r["pass"] == "是")
    fails = [r["id"] for r in rows if r["pass"] != "是"]
    ws3 = wb.create_sheet("03_汇总")
    ws3.append(["用例数", "达标", "未达标", "结论"])
    ws3.append([len(rows), passed, ",".join(fails) or "无", "用户论据×输出质量联测"])
    for c in range(1, 5):
        ws3.cell(1, c).fill = RED
        ws3.cell(1, c).font = WHITE

    for sheet in wb.worksheets:
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = WRAP

    wb.save(WB1)
    try:
        wb.save(WB2)
        desk = WB2
    except PermissionError:
        desk = Path.home() / "Desktop" / "skill_eval_v0.4_joint_v2.xlsx"
        wb.save(desk)
    print("saved", WB1)
    print("saved", desk)
    print("pass", passed, "/", len(rows), "fails", fails)


if __name__ == "__main__":
    main()
