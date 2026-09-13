# -*- coding: utf-8 -*-
"""Score news-wire outputs on 6 dims from corpus training goals; build Excel."""

from __future__ import annotations

import csv
import json
import re
import shutil
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "eval" / "outputs_news_v040"
CASES = ROOT / "eval" / "skill_eval_cases_v0.4_news.csv"
WB = ROOT / "eval" / "skill_eval_v0.4_news.xlsx"
DESK = Path.home() / "Desktop" / "skill_eval_v0.4_news.xlsx"
CHANGELOG = ROOT / "eval" / "news_skill_changelog.md"

WRAP = Alignment(wrap_text=True, vertical="top")
HDR = PatternFill("solid", fgColor="1F4E79")
HDR_F = Font(color="FFFFFF", bold=True)
FAIL = PatternFill("solid", fgColor="F8CBAD")
PASS = PatternFill("solid", fgColor="C6EFCE")

GENZ = [
    "imagine ",
    "vibes",
    "solar-punk",
    "solarpunk",
    "low-key",
    "would you visit",
    "paradise",
    "🔥",
    "🏍️",
    "☀️",
    "🌿",
    "🐑",
    "emoji",
]
FAKE_MARKERS = [
    "100 million kilowatt",
    "1亿千瓦",
    "10 million tourists",
    "游客破千万",
    "export to europe and north america sales no.1",
    "出口欧美销量第一",
    "ranking first in exports",
]
LEXICON_POS = [
    "clean energy",
    "win-win",
    "livelihood",
    "ecological",
    "desertification",
    "rural revitalization",
    "photovoltaic",
    "forage",
    "green development",
    "清洁能源",
    "双赢",
    "生态",
    "治沙",
    "乡村振兴",
    "光伏",
    "牧光",
]


def _body(d: dict) -> str:
    return (d.get("article") or d.get("post") or "").strip()


def _score_tone(text: str, lang: str) -> tuple[int, str]:
    t = text.lower()
    notes = []
    score = 4
    if re.search(r"\b(i |we |let's |come visit)\b", t) and "xinhua" not in t[:80]:
        # first person tourism
        if t.startswith("i ") or " i " in t[:200]:
            score -= 2
            notes.append("第一人称偏多")
    if any(g in t for g in ("lol", "omg", "gonna", "kinda")):
        score -= 1
        notes.append("口语过网感")
    if "said" in t or "表示" in text or "说" in text or "recalled" in t:
        score = min(5, score + 0)
    else:
        if len(text) > 400:
            score = min(score, 3)
            notes.append("缺少转述引语痕迹")
    return max(1, min(5, score)), ";".join(notes)


def _score_narrative(text: str, wire_plan: dict) -> tuple[int, str]:
    t = text.lower()
    beats = 0
    notes = []
    if re.search(r"(once|past|for generations|曾经|过去|昔日)", t):
        beats += 1
    else:
        notes.append("缺今昔/背景")
    if re.search(r"(panel|solar|光伏|solar panel)", t):
        beats += 1
    if re.search(r"(sheep|herder|牧民|羊|graze)", t) or re.search(
        r"(workshop|training|luban|鲁班)", t
    ):
        beats += 1
    if re.search(
        r"(because|shade|agreement|evaporat|mowing|机制|协议|遮挡|蒸发)", t
    ):
        beats += 1
    else:
        notes.append("机制因果弱")
    if re.search(r"(\d+|percent|yuan|平方公里|只|万元)", t):
        beats += 1
    if wire_plan.get("angle") or wire_plan.get("lead_type"):
        beats = min(5, beats + 1)
    score = 1 + min(4, beats // 1)
    score = min(5, max(1, 1 + beats))
    if beats <= 2:
        score = min(score, 2)
    elif beats == 3:
        score = 3
    elif beats == 4:
        score = 4
    else:
        score = 5
    return score, ";".join(notes)


def _score_emphasis(text: str) -> tuple[int, str]:
    t = text.lower()
    flags = {
        "energy": bool(re.search(r"(clean energy|electricity|power|发电|清洁能源|绿电)", t)),
        "ecology": bool(
            re.search(r"(ecological|desert|vegetation|grass|治沙|生态|植被|荒漠)", t)
        ),
        "livelihood": bool(
            re.search(r"(herder|income|livelihood|forage|增收|牧民|收入|就业|skill)", t)
        ),
    }
    n = sum(flags.values())
    notes = [k for k, v in flags.items() if v]
    if n >= 3:
        return 5, "+".join(notes)
    if n == 2:
        return 4, "+".join(notes)
    if n == 1:
        return 2, "+".join(notes) or "单点"
    return 1, "无强调面"


def _score_lexicon(text: str) -> tuple[int, str]:
    t = text.lower()
    pos = sum(1 for w in LEXICON_POS if w in t or w in text)
    score = 3
    if pos >= 3:
        score = 5
    elif pos == 2:
        score = 4
    elif pos == 1:
        score = 3
    else:
        score = 2
    bad = [g for g in GENZ if g in t]
    if bad:
        score = max(1, score - 2)
        return score, f"口径词{pos};禁用:{','.join(bad[:3])}"
    return score, f"口径词命中约{pos}"


def _score_genre_sep(data: dict, text: str) -> tuple[int, str]:
    t = text.lower()
    score = 5
    notes = []
    if data.get("genre") != "news":
        score -= 2
        notes.append(f"genre={data.get('genre')}")
    skills = data.get("skills_applied") or []
    if "china_story_news" not in skills and data.get("genre") == "news":
        score -= 1
        notes.append("未注入news skill")
    if "china_story_post" in skills and data.get("genre") == "news":
        # shouldn't be for news route
        notes.append("含post skill")
    for g in GENZ:
        if g in t:
            score -= 1
            notes.append(g.strip())
            break
    if re.search(r"#[A-Za-z]\w+", text) and len(re.findall(r"#\w+", text)) >= 3:
        score -= 1
        notes.append("hashtag堆砌")
    if (data.get("five_w") or {}) and data.get("genre") == "news":
        # empty five_w is ok; non-empty weird but not fail
        pass
    return max(1, min(5, score)), ";".join(notes)


def _score_evidence(data: dict, text: str, mats: str, expect_block: bool) -> tuple[int, str]:
    if expect_block:
        if data.get("error") or not text:
            return 5, "正确拦截"
        # wrote something off-topic
        if re.search(r"(manchester|man utd|曼联|控球)", text.lower()):
            return 1, "硬写成稿"
        return 2, "应拦截却成稿"

    score = 4
    notes = []
    for m in FAKE_MARKERS:
        if m.lower() in text.lower():
            score = 1
            notes.append(f"疑似编造:{m}")
            break
    # numbers in text should roughly appear in materials or evidence
    ev = " ".join(
        (e.get("statement") or "") for e in (data.get("evidence_used") or [])
    )
    pool = (mats + " " + ev).lower()
    # check a few salient invented patterns already covered
    notes_list = data.get("evidence_notes") or []
    if text and len(text) > 300 and not notes_list:
        score = max(1, score - 1)
        notes.append("缺evidence_notes")
    if "100,000" in text or "100000" in text.replace(",", ""):
        if "100,000" not in pool and "100000" not in pool.replace(",", "") and "10万" not in mats:
            score = max(1, score - 2)
            notes.append("收入数字来源可疑")
    return max(1, min(5, score)), ";".join(notes)


def score_one(data: dict, case: dict) -> dict:
    meta = data.get("_meta") or {}
    cid = meta.get("id") or case.get("测试ID")
    text = _body(data)
    expect_block = bool(meta.get("expect_block") or cid == "N11")
    mats = case.get("用户资料（粘贴）") or ""

    if expect_block:
        d_ev, n_ev = _score_evidence(data, text, mats, True)
        # other dims: if blocked successfully, give neutral-high on structure dims
        blocked = bool(data.get("error")) or not text
        base = 5 if blocked else 2
        return {
            "id": cid,
            "style": meta.get("prompt_style") or case.get("提示词形态"),
            "genre": data.get("genre"),
            "d_tone": base,
            "d_narrative": base,
            "d_emphasis": base,
            "d_lexicon": base,
            "d_genre_sep": 5 if blocked else 2,
            "d_evidence": d_ev,
            "total": base * 5 + d_ev if blocked else 2 * 5 + d_ev,
            "pass": "是" if blocked and d_ev >= 4 else "否",
            "problem": n_ev if not blocked else "门控OK",
            "excerpt": (text or str(data.get("error") or ""))[:280],
            "focus": case.get("评分重点"),
            "expect": case.get("用户要求（本条预期）"),
            "prompt": case.get("用户输入Prompt（完整）"),
            "skills": ",".join(data.get("skills_applied") or []),
            "notes_detail": {
                "tone": "gate",
                "narrative": "gate",
                "emphasis": "gate",
                "lexicon": "gate",
                "genre_sep": "gate",
                "evidence": n_ev,
            },
        }

    if data.get("error") and not text:
        return {
            "id": cid,
            "style": meta.get("prompt_style"),
            "genre": data.get("genre"),
            "d_tone": 1,
            "d_narrative": 1,
            "d_emphasis": 1,
            "d_lexicon": 1,
            "d_genre_sep": 1,
            "d_evidence": 1,
            "total": 6,
            "pass": "否",
            "problem": f"未成稿:{str(data.get('error'))[:120]}",
            "excerpt": "",
            "focus": case.get("评分重点"),
            "expect": case.get("用户要求（本条预期）"),
            "prompt": case.get("用户输入Prompt（完整）"),
            "skills": ",".join(data.get("skills_applied") or []),
            "notes_detail": {},
        }

    wire = data.get("wire_plan") or {}
    d1, n1 = _score_tone(text, case.get("language") or "English")
    d2, n2 = _score_narrative(text, wire)
    d3, n3 = _score_emphasis(text)
    d4, n4 = _score_lexicon(text)
    d5, n5 = _score_genre_sep(data, text)
    d6, n6 = _score_evidence(data, text, mats, False)
    total = d1 + d2 + d3 + d4 + d5 + d6
    problems = [x for x in [n1, n2, n4, n5, n6] if x]
    # N06 special: must not contain genz
    if cid == "N06" and any(g in text.lower() for g in GENZ):
        d5 = 1
        total = d1 + d2 + d3 + d4 + d5 + d6
        problems.append("N06串味")
    if cid == "N07" and any(m.lower() in text.lower() for m in FAKE_MARKERS):
        d6 = 1
        total = d1 + d2 + d3 + d4 + d5 + d6
        problems.append("N07编造")

    return {
        "id": cid,
        "style": meta.get("prompt_style") or case.get("提示词形态"),
        "genre": data.get("genre"),
        "d_tone": d1,
        "d_narrative": d2,
        "d_emphasis": d3,
        "d_lexicon": d4,
        "d_genre_sep": d5,
        "d_evidence": d6,
        "total": total,
        "pass": "是" if total >= 22 and d5 >= 3 and d6 >= 3 else "否",
        "problem": "; ".join(problems)[:200],
        "excerpt": text[:320],
        "focus": case.get("评分重点"),
        "expect": case.get("用户要求（本条预期）"),
        "prompt": case.get("用户输入Prompt（完整）"),
        "skills": ",".join(data.get("skills_applied") or []),
        "notes_detail": {
            "tone": n1,
            "narrative": n2,
            "emphasis": n3,
            "lexicon": n4,
            "genre_sep": n5,
            "evidence": n6,
        },
    }


def build_workbook(rows: list[dict], changelog: str):
    wb = Workbook()

    # 01 detail
    ws = wb.active
    ws.title = "01_评测明细"
    headers = [
        "测试ID",
        "提示词形态",
        "genre",
        "语气",
        "讲述",
        "强调面",
        "用词口径",
        "与贴文隔离",
        "证据约束",
        "总分/30",
        "达标",
        "问题摘要",
        "skills",
        "成稿摘录",
    ]
    ws.append(headers)
    for c in ws[1]:
        c.fill = HDR
        c.font = HDR_F
    for r in rows:
        ws.append(
            [
                r["id"],
                r.get("style"),
                r.get("genre"),
                r["d_tone"],
                r["d_narrative"],
                r["d_emphasis"],
                r["d_lexicon"],
                r["d_genre_sep"],
                r["d_evidence"],
                r["total"],
                r["pass"],
                r.get("problem"),
                r.get("skills"),
                r.get("excerpt"),
            ]
        )
        fill = PASS if r["pass"] == "是" else FAIL
        ws.cell(ws.max_row, 11).fill = fill

    # 02 cases
    ws2 = wb.create_sheet("02_用例与Prompt")
    ws2.append(["测试ID", "完整Prompt", "预期", "评分重点", "资料字数"])
    for c in ws2[1]:
        c.fill = HDR
        c.font = HDR_F
    cases = list(csv.DictReader(CASES.open(encoding="utf-8-sig")))
    case_map = {c["测试ID"]: c for c in cases}
    for r in rows:
        c = case_map.get(r["id"], {})
        ws2.append(
            [
                r["id"],
                c.get("用户输入Prompt（完整）"),
                c.get("用户要求（本条预期）"),
                c.get("评分重点"),
                len(c.get("用户资料（粘贴）") or ""),
            ]
        )

    # 03 dimensions guide
    ws3 = wb.create_sheet("03_六维说明")
    ws3.append(["维度", "对应老师/语料目标", "1分", "5分"])
    for c in ws3[1]:
        c.fill = HDR
        c.font = HDR_F
    guide = [
        ("语气", "官媒第三人称冷静乐观", "博主/嘲讽/过网感", "通稿转述+克制"),
        ("讲述", "导语→人→机制→数据→意义", "无因果碎片", "节拍齐全可核对"),
        ("强调面", "生态/能源/民生(证据内)", "只喊口号", "2–3面落地"),
        ("用词口径", "clean energy/win-win等", "solar-punk/vibes", "官方词典+无禁用"),
        ("与贴文隔离", "独立news skill/禁5W短帖腔", "emoji/Imagine/visit", "genre=news且无串味"),
        ("证据约束", "数字引语仅evidence", "编造装机游客销量", "可映射且notes齐全"),
    ]
    for g in guide:
        ws3.append(list(g))

    # 04 changelog
    ws4 = wb.create_sheet("04_Skill改动")
    ws4.append(["内容"])
    ws4["A1"].fill = HDR
    ws4["A1"].font = HDR_F
    for i, line in enumerate((changelog or "").splitlines(), start=2):
        ws4.cell(i, 1, line)

    # 05 summary
    ws5 = wb.create_sheet("05_汇总看板")
    n = len(rows)
    passed = sum(1 for r in rows if r["pass"] == "是")
    avg = sum(r["total"] for r in rows) / n if n else 0
    dims = ["d_tone", "d_narrative", "d_emphasis", "d_lexicon", "d_genre_sep", "d_evidence"]
    labels = ["语气", "讲述", "强调面", "用词口径", "与贴文隔离", "证据约束"]
    ws5.append(["指标", "数值"])
    ws5.append(["样本数", n])
    ws5.append(["达标数", passed])
    ws5.append(["达标率", f"{passed}/{n}"])
    ws5.append(["平均分/30", round(avg, 2)])
    for lab, d in zip(labels, dims):
        ws5.append([f"均分-{lab}", round(sum(r[d] for r in rows) / n, 2) if n else 0])
    fails = [r["id"] for r in rows if r["pass"] != "是"]
    ws5.append(["未达标ID", ", ".join(fails) or "无"])

    for wsx in wb.worksheets:
        for col in range(1, wsx.max_column + 1):
            wsx.column_dimensions[get_column_letter(col)].width = 18 if col < 12 else 40
        for row in wsx.iter_rows():
            for cell in row:
                cell.alignment = WRAP

    wb.save(WB)
    try:
        shutil.copyfile(WB, DESK)
    except Exception:
        pass
    return WB, DESK, avg, passed, n


def main():
    cases = list(csv.DictReader(CASES.open(encoding="utf-8-sig")))
    case_map = {c["测试ID"]: c for c in cases}
    rows = []
    for path in sorted(OUT.glob("N*.json")):
        data = json.loads(path.read_text(encoding="utf-8"))
        cid = (data.get("_meta") or {}).get("id") or path.stem
        rows.append(score_one(data, case_map.get(cid, {"测试ID": cid})))
    rows.sort(key=lambda x: x["id"])
    changelog = CHANGELOG.read_text(encoding="utf-8") if CHANGELOG.exists() else ""
    wb, desk, avg, passed, n = build_workbook(rows, changelog)
    summary = {
        "n": n,
        "passed": passed,
        "avg": avg,
        "rows": rows,
        "workbook": str(wb),
        "desktop": str(desk),
    }
    (OUT / "_scores.json").write_text(
        json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print(f"scored {n} avg={avg:.2f} pass={passed}/{n}")
    print("workbook", wb)
    for r in rows:
        print(
            r["id"],
            r["total"],
            r["pass"],
            r["d_tone"],
            r["d_narrative"],
            r["d_emphasis"],
            r["d_lexicon"],
            r["d_genre_sep"],
            r["d_evidence"],
            r.get("problem", "")[:60],
        )


if __name__ == "__main__":
    main()
