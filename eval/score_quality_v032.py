"""Score v0.3.2 quality outputs: 5W / audience / authority (not gate-centric)."""

from __future__ import annotations

import json
import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "eval" / "outputs_quality_v032"
WB1 = ROOT / "eval" / "skill_eval_v0.3.2_quality.xlsx"
WB2 = Path.home() / "Desktop" / "skill_eval_v0.3.2_quality.xlsx"

RED = PatternFill("solid", fgColor="C41E3A")
WHITE = Font(color="FFFFFF", bold=True)
WRAP = Alignment(wrap_text=True, vertical="top")

SHELL_5W = re.compile(r"^(who|what|audience|platform|channel|effect|n/?a|-|—|\.{0,3})$", re.I)
ATTACK = ["碾压", "糟糕", "loser", "food coma", "don't get culture", "不懂文化", "adorable"]


def _blob(d: dict) -> str:
    return json.dumps(d, ensure_ascii=False).lower()


def score_one(data: dict) -> dict:
    meta = data.get("_meta") or {}
    cid = meta.get("id", "?")
    if data.get("error") and not (data.get("post") or "").strip():
        return {
            "id": cid,
            "style": meta.get("prompt_style"),
            "q_5w": 1,
            "q_audience": 1,
            "q_authority": 1,
            "q_total": 3,
            "pass": "否",
            "problem": f"未成稿: {str(data.get('error'))[:100]}",
            "prompt": meta.get("full_prompt"),
            "expect": meta.get("expect"),
            "focus": meta.get("focus"),
            "excerpt": "",
            "five_w": "",
        }

    fw = data.get("five_w") or data.get("5w") or {}
    post = (data.get("post") or "").strip()
    notes = " ".join(data.get("evidence_notes") or []).lower()
    ev = data.get("evidence_used") or []
    ev_text = " ".join((e.get("statement") or "") + " " + (e.get("source") or "") for e in ev).lower()
    country = (data.get("country") or meta.get("expect") or "").lower()
    # country from result
    country = (data.get("country") or "").lower()

    # --- Q-5W (1-5) ---
    keys = ["who", "says_what", "channel", "to_whom", "effect"]
    filled = 0
    for k in keys:
        v = str(fw.get(k) or "").strip()
        if v and not SHELL_5W.match(v) and len(v) >= 8:
            filled += 1
    q5 = 1 + filled  # 1..6 -> clamp 5
    q5 = min(5, q5)
    # consistency: says_what token overlap with post
    sw = str(fw.get("says_what") or "").lower()
    if sw and post:
        tokens = [t for t in re.split(r"\W+", sw) if len(t) > 3][:6]
        if tokens and sum(1 for t in tokens if t in post.lower()) >= max(1, len(tokens) // 3):
            q5 = min(5, q5 + 0)  # keep
        elif tokens and not any(t in post.lower() for t in tokens):
            q5 = max(1, q5 - 1)

    # --- Q-audience (1-5) ---
    tw = str(fw.get("to_whom") or "").lower()
    post_l = post.lower()
    qa = 3
    country_markers = {
        "america": ["america", "american", "us ", "u.s", "美国", "weekend", "road trip", "feed"],
        "uk": ["uk", "britain", "british", "london", "英国", "tea", "weekend market"],
        "japan": ["japan", "japanese", "日本", "season"],
    }
    # detect from meta prompt
    prompt = (meta.get("full_prompt") or "").lower()
    target = "america"
    if "uk" in prompt or "英国" in prompt:
        target = "uk"
    elif "japan" in prompt or "日本" in prompt:
        target = "japan"
    elif "america" in prompt or "美国" in prompt or "us audience" in prompt:
        target = "america"

    markers = country_markers.get(target, country_markers["america"])
    hit_tw = any(m in tw for m in markers + [target])
    hit_post = any(m in post_l for m in markers)
    if hit_tw and hit_post:
        qa = 5
    elif hit_tw or hit_post:
        qa = 4
    elif len(tw) >= 12:
        qa = 3
    else:
        qa = 2
    if any(a in post_l for a in ATTACK):
        qa = min(qa, 2)

    # --- Q-authority (1-5) ---
    qauth = 3
    needs_unesco = "unesco" in prompt or "非遗" in prompt or "春节" in (meta.get("parsed_theme") or "")
    has_unesco_ev = "unesco" in ev_text or "intangible" in ev_text or "非遗" in ev_text
    has_unesco_out = "unesco" in post_l or "unesco" in notes or "intangible" in post_l
    if needs_unesco and has_unesco_ev:
        qauth = 5 if has_unesco_out else 2
    elif ev:
        # notes mapping
        if "evidence#" in notes or "evidence #" in notes or re.search(r"evidence\s*#?\s*\d", notes):
            qauth = 5
        elif notes.strip():
            qauth = 4
        else:
            qauth = 3
    else:
        qauth = 2

    # fabricated vibe
    if any(x in post_l for x in ["97%", "oat milk", "zoom meeting", "5000万"]):
        qauth = min(qauth, 2)

    total = q5 + qa + qauth  # /15
    ok = total >= 11 and q5 >= 3 and qa >= 3 and qauth >= 3
    problems = []
    if q5 < 4:
        problems.append("5W偏空壳或与正文不一致")
    if qa < 4:
        problems.append("受众桥梁弱")
    if qauth < 4:
        problems.append("权威映射弱")
    if not problems:
        problems.append("质量维度达标")

    return {
        "id": cid,
        "style": meta.get("prompt_style"),
        "pair": meta.get("pair") or "",
        "q_5w": q5,
        "q_audience": qa,
        "q_authority": qauth,
        "q_total": total,
        "pass": "是" if ok else "否",
        "problem": "；".join(problems),
        "prompt": meta.get("full_prompt"),
        "expect": meta.get("expect"),
        "focus": meta.get("focus"),
        "excerpt": post[:220],
        "five_w": json.dumps(fw, ensure_ascii=False)[:280],
    }


def main():
    import csv

    case_pair = {}
    case_csv = ROOT / "eval" / "skill_eval_cases_v0.3.2_quality.csv"
    if case_csv.exists():
        for r in csv.DictReader(case_csv.open(encoding="utf-8-sig")):
            case_pair[r["测试ID"]] = r.get("稳定配对") or ""

    rows = []
    for p in sorted(OUT.glob("Q*.json")):
        data = json.loads(p.read_text(encoding="utf-8"))
        meta = data.setdefault("_meta", {})
        if not meta.get("pair"):
            meta["pair"] = case_pair.get(meta.get("id") or p.stem, "")
        rows.append(score_one(data))

    wb = Workbook()
    ws = wb.active
    ws.title = "01_输出质量明细"
    headers = [
        "测试ID",
        "稳定配对",
        "提示词形态",
        "提示词（完整）",
        "用户要求",
        "评分重点",
        "Q-5W/5",
        "Q-受众/5",
        "Q-权威/5",
        "合计/15",
        "是否达标",
        "主要问题",
        "5W摘要",
        "正文摘要",
    ]
    ws.append(headers)
    for r in rows:
        ws.append(
            [
                r["id"],
                r.get("pair") or "",
                r["style"],
                r["prompt"],
                r["expect"],
                r["focus"],
                r["q_5w"],
                r["q_audience"],
                r["q_authority"],
                r["q_total"],
                r["pass"],
                r["problem"],
                r["five_w"],
                r["excerpt"],
            ]
        )
    for c in range(1, len(headers) + 1):
        cell = ws.cell(1, c)
        cell.fill = RED
        cell.font = WHITE
        cell.alignment = WRAP
    for col in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(col)].width = 18
    ws.column_dimensions["C"].width = 48
    ws.column_dimensions["L"].width = 40
    ws.column_dimensions["M"].width = 40

    ws2 = wb.create_sheet("02_与上轮差别")
    ws2.append(["项目", "上一阶段（门控/表面）", "本轮 v0.3.2（输出质量）"])
    diffs = [
        ("评测焦点", "主题是否库内、能否STOP", "5W是否可执行、受众是否适配、权威是否进正文"),
        ("提示词", "多为「主题：xxx + 参数表」", "极短/疑问/口语/英文/角色长提示/简报/作业口吻 + Q25-40长短对"),
        ("Skill改动", "易停在「避免元叙事凑字」类禁令", "多样口令前已升 v0.3.2；全合格=改skill后评测，非原skill不动"),
        ("Prompt模板", "列出5W名词", "要求5W与正文一致、evidence_notes编号映射"),
        ("分数维", "D1–D7+门控", "本表新增 Q-5W / Q-受众 / Q-权威（/15）"),
        ("门控", "继续保留，但不作为本轮改进叙事", "本用例几乎全是应成稿题"),
    ]
    for d in diffs:
        ws2.append(list(d))
    for c in range(1, 4):
        ws2.cell(1, c).fill = RED
        ws2.cell(1, c).font = WHITE
    for col in range(1, 4):
        ws2.column_dimensions[get_column_letter(col)].width = 36

    # stability pairs
    ws4 = wb.create_sheet("04_长短提示稳定性")
    ws4.append(
        [
            "配对",
            "短题ID",
            "长题ID",
            "短合计",
            "长合计",
            "分差|短-长|",
            "短达标",
            "长达标",
            "稳定性",
            "说明",
        ]
    )
    by_id = {r["id"]: r for r in rows}
    pairs = {}
    for r in rows:
        pid = r.get("pair") or ""
        if not pid:
            continue
        pairs.setdefault(pid, []).append(r)
    unstable = []
    for pid, items in sorted(pairs.items()):
        short = next((x for x in items if "短" in (x.get("style") or "")), None)
        long = next((x for x in items if "长" in (x.get("style") or "")), None)
        if not short or not long:
            # fallback by id order
            items_sorted = sorted(items, key=lambda x: x["id"])
            if len(items_sorted) >= 2:
                short, long = items_sorted[0], items_sorted[1]
            else:
                continue
        diff = abs(short["q_total"] - long["q_total"])
        both = short["pass"] == "是" and long["pass"] == "是"
        stable = both and diff <= 2
        if not stable:
            unstable.append(pid)
        ws4.append(
            [
                pid,
                short["id"],
                long["id"],
                short["q_total"],
                long["q_total"],
                diff,
                short["pass"],
                long["pass"],
                "稳定" if stable else "波动",
                "同分位达标且分差≤2" if stable else "需复查短令或长约束是否塌质量",
            ]
        )
    for c in range(1, 11):
        ws4.cell(1, c).fill = RED
        ws4.cell(1, c).font = WHITE

    ws3 = wb.create_sheet("03_汇总")
    n = len(rows) or 1
    mean5 = round(sum(r["q_5w"] for r in rows) / n, 2)
    meana = round(sum(r["q_audience"] for r in rows) / n, 2)
    meanu = round(sum(r["q_authority"] for r in rows) / n, 2)
    meant = round(sum(r["q_total"] for r in rows) / n, 2)
    passed = sum(1 for r in rows if r["pass"] == "是")
    fails = [r["id"] for r in rows if r["pass"] != "是"]
    ws3.append(
        [
            "用例数",
            "达标数",
            "Q-5W均分",
            "Q-受众均分",
            "Q-权威均分",
            "合计均分/15",
            "未达标ID",
            "稳定配对波动",
            "结论",
            "Skill说明",
        ]
    )
    ws3.append(
        [
            len(rows),
            passed,
            mean5,
            meana,
            meanu,
            meant,
            ",".join(fails) or "无",
            ",".join(unstable) or "无",
            "Q扩至40：长短提示稳定性；门控非主线",
            "全合格建立在 v0.3.2 skill+prompt 之上，非原表面skill不动",
        ]
    )
    for c in range(1, 11):
        ws3.cell(1, c).fill = RED
        ws3.cell(1, c).font = WHITE

    for sheet in wb.worksheets:
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = WRAP

    wb.save(WB1)
    try:
        wb.save(WB2)
        desk = WB2
    except PermissionError:
        desk = Path.home() / "Desktop" / "skill_eval_v0.3.2_quality_v2.xlsx"
        wb.save(desk)
    print("saved", WB1)
    print("saved", desk)
    print("mean", meant, "pass", passed, "/", len(rows), "fails", fails, "unstable", unstable)


if __name__ == "__main__":
    main()
