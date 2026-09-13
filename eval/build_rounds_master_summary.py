"""短贴文阶段 · 每轮修改总表（归档后转向用户论据）。"""

from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
OUT1 = ROOT / "eval" / "skill_rounds_master_summary.xlsx"
OUT2 = Path.home() / "Desktop" / "skill_rounds_master_summary.xlsx"

RED = PatternFill("solid", fgColor="C41E3A")
WHITE = Font(color="FFFFFF", bold=True)
WRAP = Alignment(wrap_text=True, vertical="top")


def style_header(ws, n):
    for c in range(1, n + 1):
        cell = ws.cell(1, c)
        cell.fill = RED
        cell.font = WHITE
        cell.alignment = WRAP


def widen(ws, w=22):
    for col in range(1, ws.max_column + 1):
        ws.column_dimensions[get_column_letter(col)].width = w


def main():
    wb = Workbook()

    # --- Sheet1 总表 ---
    ws = wb.active
    ws.title = "00_每轮修改总表"
    headers = [
        "轮次",
        "日期",
        "阶段结论",
        "版本",
        "改动焦点",
        "改动文件",
        "权威/依据",
        "评测规模",
        "关键结果",
        "是否保留",
        "与老师/立项对齐",
    ]
    ws.append(headers)
    rows = [
        (
            1,
            "2026-08-22",
            "贴文基线",
            "v0.1",
            "建立 D1–D7 评测与 intl-comm 基线",
            "skills/intl-comm.md；eval 15题",
            "立项7指标改编",
            "15",
            "均分约26；暴露库外细节/语气问题",
            "是",
            "对齐立项主动传播主线",
        ),
        (
            2,
            "2026-08-22",
            "贴文证据+语气",
            "v0.1→v0.2",
            "证据硬性/演示期分层；语气边界；自检",
            "intl-comm.md",
            "评测失败+产品边界",
            "抽检",
            "T05改善；T11 sarcastic 仍弱",
            "是",
            "对齐「事实来自论据」",
        ),
        (
            3,
            "2026-08-22",
            "贴文门控",
            "v0.2+gate",
            "主题-素材对齐门控（拒编）",
            "kb_local.py；story_post_gen.py",
            "B组漏拦",
            "门控题",
            "T09/T10/T14 等 STOP 成功",
            "是",
            "对齐不编造；拒库外域",
        ),
        (
            4,
            "2026-09-06",
            "贴文体裁拆分",
            "v0.3-genre",
            "帖文专用 Lasswell 5W；他体裁占位",
            "genres/*；genre_router；skills.yaml",
            "Lasswell1948；老师0822先贴文",
            "40",
            "genre=post 注入成功；南方周末不硬套短帖",
            "是",
            "对齐老师「贴文用5W」",
        ),
        (
            5,
            "2026-09-06",
            "贴文语气闭环",
            "v0.3.1",
            "sarcastic 不得嘲讽受众；用户批准草案落地",
            "intl-comm§C；china-story-post；story_post_prompt",
            "用户批准；少抬杠",
            "T11/T33/T12 重测",
            "嘲讽受众消除；40题均分29.27未达标0",
            "是",
            "对齐国际传播受众尊重",
        ),
        (
            6,
            "2026-09-06",
            "表结构可读性",
            "表结构v2",
            "评测表加提示词+用户要求；删技术追踪列",
            "build_full_8sheet_workbook.py",
            "用户截图反馈",
            "表头",
            "01/02 可见完整 Prompt 与预期",
            "是",
            "评测可读性；非产品能力",
        ),
        (
            7,
            "2026-09-06",
            "【转向】用户论据",
            "v0.4-evidence 启动",
            "短贴文阶段归档；启动用户资料→evidence 双源",
            "evidence-user-materials.md；story_post_gen；API/前端粘贴",
            "老师P0上传资料；立项不编造",
            "论据向新用例（本轮起）",
            "见 01_与老师对齐说明",
            "进行中",
            "补齐此前最大缺口（见对齐说明）",
        ),
        (
            8,
            "2026-09-06",
            "【纠偏】输出质量",
            "v0.3.2-post",
            "用户指出：勿再停在门控/表面禁令；转向5W可执行+受众+权威+多样提示词",
            "china-story-post.md；story_post_prompt；Q01-Q24评测",
            "Lasswell1948；老师贴文用5W；用户反馈",
            "24条多样问法",
            "见 skill_eval_v0.3.2_quality.xlsx",
            "是",
            "对齐「内容怎么讲」而非仅「能不能进库」",
        ),
        (
            9,
            "2026-09-06",
            "长短提示稳定性",
            "v0.3.2 + Q40",
            "Q扩至40；8对短/长同角度；专打稳定性",
            "skill_eval_cases_v0.3.2_quality.csv Q25-40",
            "用户指定顺序①",
            "40",
            "40/40达标；配对分差≤2全稳定",
            "是",
            "说明：全合格基于已改v0.3.2，非原表面skill",
        ),
        (
            10,
            "2026-09-06",
            "用户论据×质量联测",
            "v0.4-joint",
            "用户上传+5W/受众/权威打通",
            "joint cases J01-J12；evidence skill+post skill",
            "老师P0+用户指定顺序②",
            "12联测",
            "J01-J12 全达标（含2条正确拒编）",
            "是",
            "补齐喂资料与讲法质量闭环",
        ),
    ]
    for r in rows:
        ws.append(list(r))
    style_header(ws, len(headers))
    widen(ws, 18)
    ws.column_dimensions["E"].width = 36
    ws.column_dimensions["F"].width = 40
    ws.column_dimensions["K"].width = 28

    # --- Sheet2 对齐说明 ---
    ws2 = wb.create_sheet("01_与老师立项对齐")
    ws2.append(["检查项", "老师/立项要求", "当前状态", "本轮决定", "风险若继续只测短贴文"])
    align = [
        (
            "短贴文 Skill",
            "先把贴文做好；贴文用5W",
            "v0.3.1 已可用；40题回归绿",
            "暂停大规模贴文扩样；保留回归抽检",
            "边际收益低，挤占 P0 时间",
        ),
        (
            "用户上传资料",
            "上传资料再生成；用户资料+本地库",
            "此前 0%；本轮启动 skill+粘贴融合",
            "优先做论据调用",
            "直接违背老师 P0",
        ),
        (
            "论据可区分",
            "面板标本地库/用户上传",
            "本轮加 source_type",
            "前后端同步标注",
            "不可审计「是否喂了资料」",
        ),
        (
            "多体裁",
            "新闻/深度/脚本可扩展",
            "占位；待南方周末材料",
            "不编造框架",
            "无材料硬写=不符合「有据」",
        ),
        (
            "知识图谱",
            "立项名称含知识图谱",
            "现为本地 JSON/Chroma RAG",
            "本阶段不假装已有图谱",
            "宣称图谱但未实现=立项名实不符（已知债）",
        ),
        (
            "人机工作流确认",
            "定主题→喂资料→口径→初稿→多轮",
            "单步生成+改稿芯片",
            "P1 稍后；先打通喂资料",
            "跳过喂资料环节不符合会议流程",
        ),
        (
            "再跑40条短帖提示词",
            "扩样完善贴文 skill",
            "主题/角度已覆盖 A/B/C",
            "不作为门控主线；改为多样问法测输出质量(Q01-Q24)",
            "只测门控=与「怎么讲好」脱节",
        ),
        (
            "输出质量（5W/受众/权威）",
            "贴文用5W；科学有据讲给特定受众",
            "v0.3.2 skill+prompt+Q评测已启动",
            "本轮主线之一",
            "若只改「少口号」表面句=用户已批评的浅改",
        ),
    ]
    for r in align:
        ws2.append(list(r))
    style_header(ws2, 5)
    widen(ws2, 28)

    # --- Sheet3 下一阶段清单 ---
    ws3 = wb.create_sheet("02_论据阶段待办")
    ws3.append(["优先级", "事项", "产出", "验收"])
    todos = [
        ("P0", "用户资料规范化 → evidence_used", "normalize_user_materials", "JSON 含 source_type=用户上传"),
        ("P0", "双源合并进成稿链路", "story_post_gen 接受 user_materials", "仅本地空但用户有料仍可成稿"),
        ("P0", "前端粘贴资料区", "Composer textarea", "生成请求带 user_materials"),
        ("P0", "论据面板区分来源", "badge 本地库/用户上传", "肉眼可分"),
        ("P0", "论据向评测用例", "eval/skill_eval_cases_v0.4_evidence.csv", "覆盖双源/仅用户/拒编"),
        ("P1", "文件上传 .txt/.md", "multipart", "非仅粘贴"),
        ("P1", "多步确认工作流", "步骤条", "喂资料后可停顿"),
        ("P2", "南方周末体裁材料到位后再启", "news/feature skill", "禁止无材料编造"),
    ]
    for r in todos:
        ws3.append(list(r))
    style_header(ws3, 4)
    widen(ws3, 32)

    for sheet in wb.worksheets:
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = WRAP

    wb.save(OUT1)
    try:
        wb.save(OUT2)
        print("saved", OUT2)
    except PermissionError:
        alt = Path.home() / "Desktop" / "skill_rounds_master_summary_v2.xlsx"
        wb.save(alt)
        print("desktop locked; saved", alt)
    print("saved", OUT1)


if __name__ == "__main__":
    main()
