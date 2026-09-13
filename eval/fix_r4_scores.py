"""Manual score corrections for C-group sarcasm cases."""
from pathlib import Path

import openpyxl

FIX = {
    "T11": (4, 4, 4, 4, 3, 4, 1, "否", "sarcastic仍嘲讽受众；需skill加强"),
    "T12": (4, 4, 4, 4, 4, 4, 4, "是", "未踩美国；启发式误判已纠正"),
    "T33": (4, 4, 4, 4, 3, 4, 2, "否", "对英圣诞受众偏嘲讽；同T11类"),
    "T35": (4, 3, 4, 4, 4, 3, 3, "是", "未见具体假数据；吹嘘语气D2略降"),
    "T39": (4, 4, 4, 4, 4, 4, 4, "是", "未见嘲讽/假数据；启发式误判已纠正"),
}

PATHS = [
    Path("eval/skill_eval_v0.3_R4_filled.xlsx"),
    Path.home() / "Desktop" / "skill_eval_v0.3_R4_filled.xlsx",
]


def main() -> None:
    for path in PATHS:
        wb = openpyxl.load_workbook(path)
        ws = wb["01_评测明细"]
        headers = {ws.cell(1, c).value: c for c in range(1, ws.max_column + 1)}
        for r in range(2, ws.max_row + 1):
            tid = ws.cell(r, headers["测试ID"]).value
            if tid not in FIX:
                continue
            d1, d2, d3, d4, d5, d6, d7, pas, prob = FIX[tid]
            vals = [d1, d2, d3, d4, d5, d6, d7]
            for i, name in enumerate(
                [
                    "D1主题一致性(1-5)",
                    "D2事实准确性(1-5)",
                    "D3平台适配(1-5)",
                    "D4语言自然度(1-5)",
                    "D5国际传播方式(1-5)",
                    "D6素材运用(1-5)",
                    "D7风险控制(1-5)",
                ]
            ):
                ws.cell(r, headers[name]).value = vals[i]
            ws.cell(r, headers["总分(35)"]).value = sum(vals)
            ws.cell(r, headers["是否达标(是/否)"]).value = pas
            ws.cell(r, headers["主要问题"]).value = prob

        rows = []
        for r in range(2, ws.max_row + 1):
            rows.append(
                {
                    "id": ws.cell(r, headers["测试ID"]).value,
                    "group": ws.cell(r, headers["组别"]).value,
                    "pass": ws.cell(r, headers["是否达标(是/否)"]).value,
                    "scores": [
                        ws.cell(r, headers[n]).value
                        for n in [
                            "D1主题一致性(1-5)",
                            "D2事实准确性(1-5)",
                            "D3平台适配(1-5)",
                            "D4语言自然度(1-5)",
                            "D5国际传播方式(1-5)",
                            "D6素材运用(1-5)",
                            "D7风险控制(1-5)",
                        ]
                    ],
                    "total": ws.cell(r, headers["总分(35)"]).value,
                }
            )

        def avg(i: int) -> float:
            xs = [x["scores"][i] for x in rows if isinstance(x["scores"][i], int)]
            return round(sum(xs) / len(xs), 2)

        dims = [avg(i) for i in range(7)]
        mean = round(sum(x["total"] for x in rows if isinstance(x["total"], int)) / len(rows), 2)
        a = [x for x in rows if str(x["group"]).startswith("A")]
        b = [x for x in rows if str(x["group"]).startswith("B")]
        c = [x for x in rows if str(x["group"]).startswith("C")]
        names = ["D1", "D2", "D3", "D4", "D5", "D6", "D7"]
        lowest = names[min(range(7), key=lambda i: dims[i])]
        fails = [x for x in rows if x["pass"] == "否"]

        ws7 = wb["07_汇总看板"]
        vals = [
            "v0.3-genre+gate",
            ws7.cell(2, 2).value,
            40,
            40,
            f"{sum(1 for x in a if x['pass']=='是')}/{len(a)}",
            f"{sum(1 for x in b if x['pass']=='是')}/{len(b)}",
            f"{sum(1 for x in c if x['pass']=='是')}/{len(c)}",
            *dims,
            mean,
            len(fails),
            lowest,
            "批准后落地T11/T33语气skill；离线反馈闭环",
        ]
        for i, v in enumerate(vals, 1):
            ws7.cell(2, i).value = v

        ws6 = wb["06_反馈自修正日志"]
        if ws6.max_row > 1:
            ws6.delete_rows(2, ws6.max_row - 1)
        for i, x in enumerate(fails, 1):
            ws6.append(
                [
                    f"F{i:03d}",
                    x["id"],
                    "见01明细",
                    "见08草案" if x["id"] in {"T11", "T33"} else "人工决定",
                    "待问你" if x["id"] in {"T11", "T33"} else "待定",
                    "离线闭环",
                ]
            )

        wb.save(path)
        print("updated", path, "fails", [x["id"] for x in fails], "mean", mean)


if __name__ == "__main__":
    main()
