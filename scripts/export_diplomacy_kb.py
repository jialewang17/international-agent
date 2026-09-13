# -*- coding: utf-8 -*-
"""Export MFA evidence Excel into compact JSON knowledge base."""
from __future__ import annotations

import json
import re
from pathlib import Path

import openpyxl

SRC = Path(r"D:\大创\1021AIGC中国故事材料") / "回复参考-主题参考的外交部发言_20230610(1).xlsx"
OUT_DIR = Path(r"D:\大创\anyclaw") / "knowledge" / "diplomacy"
OUT_DIR.mkdir(parents=True, exist_ok=True)

# Map excel topic names -> ABSA category labels used in prompts
TOPIC_ALIASES = {
    "governance": ["governance"],
    "sports": ["sports"],
    "genocide": ["genocide"],
    "chinese religion": ["culture", "Muslim region"],
    "forced labor": ["forced labor", "manufactory"],
    "women and children rights": ["women and children rights"],
    "human rights": ["women and children rights", "governance"],
    "muslim religion": ["Muslim region", "culture"],
    "counter-terrorism": ["governance", "concentration camp"],
    "false propaganda": ["governance", "genocide"],
    "international trade": ["manufactory", "forced labor"],
}


def shorten(text: str, limit: int = 220) -> str:
    text = re.sub(r"\s+", " ", (text or "").strip())
    if len(text) <= limit:
        return text
    return text[: limit - 1].rstrip() + "…"


def main() -> None:
    wb = openpyxl.load_workbook(SRC, data_only=True)
    ws = wb[wb.sheetnames[0]]

    by_topic: dict[str, list[dict]] = {}
    current = None
    for row in ws.iter_rows(min_row=2, values_only=True):
        seq, topic, link, point = (list(row) + [None, None, None, None])[:4]
        if topic and str(topic).strip() and not str(topic).startswith("专题"):
            current = str(topic).strip()
            by_topic.setdefault(current, [])
        if current and point and str(point).strip():
            by_topic[current].append(
                {
                    "statement": shorten(str(point)),
                    "source": (str(link).strip() if link else ""),
                }
            )

    # Build category -> statements (dedupe, max 3 each)
    category_map: dict[str, list[dict]] = {}
    for topic, items in by_topic.items():
        aliases = TOPIC_ALIASES.get(topic.lower(), [topic.lower()])
        for cat in aliases:
            bucket = category_map.setdefault(cat, [])
            for it in items:
                if len(bucket) >= 3:
                    break
                if it["statement"] not in {b["statement"] for b in bucket}:
                    bucket.append(it)

    # Ensure ABSA categories exist even if empty
    for cat in [
        "celebrity",
        "governance",
        "genocide",
        "forced labor",
        "concentration camp",
        "Chinese region",
        "Muslim region",
        "Western region",
        "manufactory",
        "culture",
        "food",
        "language",
        "women and children rights",
        "sports",
    ]:
        category_map.setdefault(cat, [])

    out_json = OUT_DIR / "evidence.json"
    out_json.write_text(
        json.dumps(
            {
                "topics_raw": {k: v[:5] for k, v in by_topic.items()},
                "by_category": category_map,
            },
            ensure_ascii=False,
            indent=2,
        ),
        encoding="utf-8",
    )

    # Also write short md per category for human browsing
    for cat, items in category_map.items():
        safe = cat.replace(" ", "_")
        md = OUT_DIR / f"{safe}.md"
        lines = [f"# {cat}", ""]
        if not items:
            lines.append("_暂无精简论据，回复时以通用中国立场事实为准。_")
        for i, it in enumerate(items, 1):
            lines.append(f"{i}. {it['statement']}")
            if it.get("source"):
                lines.append(f"   - 来源: {it['source']}")
        md.write_text("\n".join(lines) + "\n", encoding="utf-8")

    print(f"wrote {out_json}")
    print("categories:", {k: len(v) for k, v in category_map.items() if v})


if __name__ == "__main__":
    main()
